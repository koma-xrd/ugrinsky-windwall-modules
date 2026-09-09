"""Check the generated maker document at its OOXML delivery boundary.

Run with bundled document Python via ``python -m unittest discover -s tests
-p test_manual_document.py``. Cached CAD drawings avoid a CAD runtime dependency.
"""

from importlib import import_module
from pathlib import Path
import sys
import unittest
from uuid import uuid4

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'src'))


class ManualDocumentTests(unittest.TestCase):
    """Missing content, drawings, layout or accessibility must fail delivery."""

    document = None
    target = None

    def setUp(self):
        self.assertTrue(
            (ROOT / 'scripts/manual/build_manual.py').is_file(),
            'The German manual builder has not been implemented',
        )
        if self.__class__.document is None:
            # Windows sandbox ACLs reject tempfile's mode-0700 directories.
            # Use a unique artifact in the existing writable build directory.
            target = ROOT / 'build' / f'manual-test-{uuid4().hex}.docx'
            self.__class__.target = target
            builder = import_module('scripts.manual.build_manual').build_manual
            self.assertEqual(builder(ROOT, target), target)
            self.__class__.document = Document(target)
        self.doc = self.document
        self.text = '\n'.join(self.doc.element.body.itertext())

    @classmethod
    def tearDownClass(cls):
        if cls.target:
            cls.target.unlink(missing_ok=True)

    def test_required_instructions_and_safety_are_present(self):
        required = (
            'Prototyp', 'Sicherheit', 'Hauptmaße', 'Druckteile', 'Kaufteile',
            'Bambu Lab P2S', '0,4 mm', 'PLA', 'ASA', 'Coupon',
            'Generator mechanisch montieren', 'Magnete', 'Trockenlayout',
            'Testwicklung', 'Serpentinen', 'A1', 'A2', 'Emaille',
            '20', '40', '80', '0,18 mm', 'Leerlauf', 'definierte Last',
            'Widerstand', 'Temperatur', 'gleicher Drehzahl',
            'N_final = N_test * V_ac_target / V_ac_test', 'aufrunden',
            'Sieben Stufen', '-18°', 'CCW', 'M8', 'Rundlauf', 'Balance',
            'Inbetriebnahme', 'Wartung', 'Fehlersuche', 'Messblatt',
            'Zeichnungsindex', 'Quellen', '69,64', '487,84', '559,64',
            '+60°', '-60°', '10,8', '11,0', '11,2', '1,5 mm',
            'Generator nicht direkt mit dem Akku verbinden.',
            'Wasserstoff', 'Überdrehzahl', 'Feuchtigkeit', 'unbeaufsichtigt',
            'Nach Messung auswählen',
        )
        for phrase in required:
            with self.subTest(phrase=phrase):
                self.assertIn(phrase, self.text)

    def test_boms_include_all_ids_and_printable_stl_references(self):
        table_text = '\n'.join(cell.text for t in self.doc.tables
                               for row in t.rows for cell in row.cells)
        for prefix, count in (('P', 5), ('H', 17)):
            for number in range(1, count + 1):
                self.assertIn(f'{prefix}{number:02}', table_text)
        for filename in ('base_rotor_module.stl', 'standard_rotor_module.stl',
                         'top_rotor_module.stl', 'top_closure.stl',
                         'lower_magnet_rotor.stl'):
            self.assertIn(filename, table_text)

    def test_all_drawings_are_inline_accessible_and_captioned(self):
        self.assertGreaterEqual(len(self.doc.inline_shapes), 11)
        self.assertFalse(self.doc.element.body.xpath('.//wp:anchor'))
        for image in self.doc.inline_shapes:
            self.assertTrue(image._inline.docPr.get('descr', '').strip())
        captions = '\n'.join(p.text for p in self.doc.paragraphs
                             if p.style.name == 'Caption')
        for number in range(1, 12):
            self.assertIn(f'E{number:02}', captions)

    def test_a4_portrait_and_eighteen_mm_side_margins(self):
        for section in self.doc.sections:
            for actual, expected in ((section.page_width, Mm(210)),
                                     (section.page_height, Mm(297)),
                                     (section.left_margin, Mm(18)),
                                     (section.right_margin, Mm(18))):
                self.assertAlmostEqual(actual, expected, delta=700)

    def test_headings_use_continuous_black_word_styles(self):
        levels = [int(p.style.name[-1]) for p in self.doc.paragraphs
                  if p.style.name.startswith('Heading ')]
        self.assertIn(1, levels)
        self.assertIn(2, levels)
        previous = 0
        for level in levels:
            self.assertLessEqual(level, previous + 1)
            previous = level
        self.assertTrue(any(p.style.name == 'Title' for p in self.doc.paragraphs))
        for name in ('Title', 'Heading 1', 'Heading 2'):
            self.assertEqual(str(self.doc.styles[name].font.color.rgb), '000000')

    def test_tables_have_borders_repeating_headers_and_no_fixed_height(self):
        self.assertGreaterEqual(len(self.doc.tables), 4)
        for table in self.doc.tables:
            self.assertIsNotNone(table.rows[0]._tr.find('w:trPr/w:tblHeader', table._element.nsmap))
            self.assertTrue(table._tbl.xpath('./w:tblPr/w:tblBorders'))
            for height in table._tbl.xpath('.//w:trHeight'):
                self.assertNotEqual(height.get(qn('w:hRule')), 'exact')


if __name__ == '__main__':
    unittest.main()
