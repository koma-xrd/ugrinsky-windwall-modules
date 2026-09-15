"""Translation coverage and package integrity; author only in document runtime."""

import hashlib
import importlib
import json
from pathlib import Path
import re
import sys
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from tests.support import temporary_build_directory

ROOT = Path(__file__).resolve().parents[1]
DOCUMENT_RUNTIME = "codex-primary-runtime" in Path(sys.executable).parts
LOCALES = {
    "en-GB": "English", "zh-CN": "Chinese-Simplified", "hi-IN": "Hindi",
    "es-ES": "Spanish", "fr-FR": "French",
}


@unittest.skipUnless(DOCUMENT_RUNTIME, "DOCX authoring requires bundled document Python")
class V5TranslationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.module_path = ROOT / "scripts/manual/localize_manual.py"
        if not cls.module_path.exists():
            return  # RED never authors before the artifact marker.
        cls.module = importlib.import_module("scripts.manual.localize_manual")
        cls.folder = cls.enterClassContext(temporary_build_directory())
        cls.outputs = cls.module.build_translations(ROOT, cls.folder)

    def setUp(self):
        self.assertTrue(self.module_path.exists(), "Shared V5 localization builder is absent")

    def test_exact_languages_and_complete_package_structure(self):
        self.assertEqual(set(self.outputs), set(LOCALES))
        for locale, path in self.outputs.items():
            with self.subTest(locale=locale):
                self.assertEqual(path.name, f"Ugrinsky-Wind-Wall-V5-Manual-{LOCALES[locale]}.docx")
                report = self.module.audit_translation(ROOT, path, locale)
                self.assertTrue(report["structural_checks_passed"])
                self.assertEqual(report["chapter_count"], 13)
                self.assertEqual(report["figure_count"], 16)
                self.assertEqual(report["table_count"], 11)
                self.assertEqual(report["embedded_images_match_release"], 16)
                self.assertTrue(report["canonical_zip_verified"])
                self.assertFalse(report["physical_validation_verified"])

    def test_rebuild_is_byte_identical_and_german_is_unchanged(self):
        for locale, path in self.outputs.items():
            released = ROOT / "release/v5/docs" / path.name
            self.assertEqual(path.read_bytes(), released.read_bytes(), locale)
            other = self.folder / "again" / path.name
            self.module.build_translation(ROOT, locale, other)
            self.assertEqual(other.read_bytes(), path.read_bytes())
        german = ROOT / self.module.SOURCE_MANUAL
        self.assertEqual(hashlib.sha256(german.read_bytes()).hexdigest(),
                         "b9a33829e9a4e664ebf611b4a251bd244ee16e5f433f1759c0a04bd78eaa0844")

    def test_coverage_is_exact_and_translations_preserve_numeric_invariants(self):
        for locale in LOCALES:
            catalog = self.module.load_catalog(ROOT, locale)
            source = self.module.source_segments(ROOT)
            self.assertEqual(set(catalog["translations"]), set(source))
            self.module.validate_catalog(source, catalog)
            for key, translated in catalog["translations"].items():
                canonical_english_drawing = bool(re.match(r"^E\d{2}[ .]", source[key].strip()))
                if (source[key] not in self.module.SHARED_LABELS
                        and not (locale == "en-GB" and canonical_english_drawing)):
                    self.assertNotEqual(translated.strip(), source[key].strip(), key)

    def test_locales_accessibility_styles_and_critical_instruction_order(self):
        from docx import Document
        for locale, path in self.outputs.items():
            with self.subTest(locale=locale):
                doc = Document(path)
                catalog = self.module.load_catalog(ROOT, locale)["translations"]
                text = "\n".join(p.text for p in doc.paragraphs)
                self.assertEqual(doc.core_properties.language, locale)
                self.assertEqual(doc.core_properties.title, doc.paragraphs[0].text)
                self.assertEqual(doc.paragraphs[0].style.name, "Title")
                for section in doc.sections:
                    self.assertAlmostEqual(section.page_width.mm, 210, places=1)
                    self.assertAlmostEqual(section.page_height.mm, 297, places=1)
                    for side in ("top", "right", "bottom", "left"):
                        self.assertAlmostEqual(getattr(section, side + "_margin").mm, 18, places=1)
                for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
                    self.assertEqual(str(doc.styles[name].font.color.rgb), "000000")
                for table in doc.tables:
                    self.assertTrue(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
                assembly = text.split(catalog["s132"])[1].split(catalog["s169"])[0]
                order = [assembly.index(catalog[key]) for key in ("s134", "s136", "s138", "s140", "s142")]
                self.assertEqual(order, sorted(order))
                captions = [p.text for p in doc.paragraphs if p.style.name == "Caption"]
                self.assertEqual(len(captions), 16)
                self.assertEqual(doc.inline_shapes[0]._inline.docPr.get("title"), "HERO")
                self.assertEqual(doc.inline_shapes[0]._inline.docPr.get("descr"),
                                 self.module.load_catalog(ROOT, locale)["hero"]["alt_text"])
                self.assertIn(self.module.HERO_QUALIFIERS[locale],
                              doc.inline_shapes[0]._inline.docPr.get("descr"))
                self.assertEqual(captions[0], self.module.load_catalog(ROOT, locale)["hero"]["caption"])
                allowed_descriptions = set(catalog.values()) | {
                    self.module.load_catalog(ROOT, locale)["hero"]["alt_text"]}
                for shape in doc.inline_shapes:
                    description = shape._inline.docPr.get("descr")
                    self.assertGreater(len(description), 40)
                    self.assertIn(description, allowed_descriptions)
                self.assertIn("N S N S", text)
                for critical in ("s006", "s008", "s104", "s108", "s109", "s121", "s130", "s148", "s157", "s158", "s173", "s175", "s177", "s191", "s233", "s234", "s295"):
                    self.assertIn(catalog[critical], text)
                matrix = doc.tables[6]
                self.assertEqual([r.cells[1].text for r in matrix.rows[1:]], ["20", "40", "80"])
                bom = doc.tables[1]
                self.assertEqual([r.cells[0].text for r in bom.rows[1:]], ["1", "5", "1", "1", "1", "1", "1", "1"])
                if locale in ("hi-IN", "zh-CN"):
                    pattern = r"[\u0900-\u097f]" if locale == "hi-IN" else r"[\u4e00-\u9fff]"
                    self.assertTrue(all(re.search(pattern, p) for p in captions))

    def test_untranslated_instruction_and_missing_safety_phrase_are_rejected(self):
        source = self.module.source_segments(ROOT)
        catalog = self.module.load_catalog(ROOT, "en-GB")
        catalog["translations"]["s006"] = source["s006"]
        with self.assertRaisesRegex(ValueError, "German"):
            self.module.validate_catalog(source, catalog)
        catalog = self.module.load_catalog(ROOT, "en-GB")
        catalog["translations"]["s008"] = catalog["translations"]["s008"].replace("Do not connect", "Connect")
        with self.assertRaisesRegex(ValueError, "safety"):
            self.module.validate_catalog(source, catalog)

    def test_all_word_style_parts_use_the_target_language(self):
        from lxml import etree
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        for locale, path in self.outputs.items():
            with ZipFile(path) as archive:
                for name in ("word/styles.xml", "word/stylesWithEffects.xml"):
                    styles = etree.fromstring(archive.read(name))
                    self.assertEqual(set(styles.xpath("//w:lang/@w:val", namespaces=namespace)), {locale})
                    for style in ("Title", "Subtitle", "Heading1", "Heading2", "Heading3"):
                        self.assertEqual(styles.xpath(f"//w:style[@w:styleId='{style}']/w:rPr/w:color/@w:val",
                                                      namespaces=namespace), ["000000"])
    def test_missing_translation_and_changed_measurement_are_rejected(self):
        source = self.module.source_segments(ROOT)
        catalog = self.module.load_catalog(ROOT, "en-GB")
        first = next(iter(catalog["translations"]))
        catalog["translations"].pop(first)
        with self.assertRaisesRegex(ValueError, "coverage"):
            self.module.validate_catalog(source, catalog)
        catalog = self.module.load_catalog(ROOT, "en-GB")
        key = next(k for k, text in source.items() if "0,35 mm" in text)
        catalog["translations"][key] = catalog["translations"][key].replace("0.35", "0.45")
        with self.assertRaisesRegex(ValueError, "numeric"):
            self.module.validate_catalog(source, catalog)

    def test_tampered_package_is_rejected_by_read_only_audit(self):
        locale, source_path = next(iter(self.outputs.items()))
        target = self.folder / "tampered.docx"
        with ZipFile(source_path) as src, ZipFile(target, "w") as dst:
            for item in src.infolist():
                payload = src.read(item.filename)
                if item.filename == "word/document.xml":
                    payload = payload.replace(b"48-V", b"24-V")
                dst.writestr(item, payload)
        with self.assertRaisesRegex(ValueError, "translation|text|content"):
            self.module.audit_translation(ROOT, target, locale)

    def test_changed_hero_is_rejected_before_translation_is_written(self):
        target = self.folder / "must-not-exist.docx"
        hero = (ROOT / "release/v5/media/windwall-fence-hero.png").read_bytes()
        original_digest = self.module.digest
        with patch.object(self.module, "digest",
                          side_effect=lambda payload: "0" * 64 if payload == hero else original_digest(payload)):
            with self.assertRaisesRegex(ValueError, "hero"):
                self.module.build_translation(ROOT, "en-GB", target)
        self.assertFalse(target.exists())

    def test_release_audit_tracks_every_catalog_and_document(self):
        audit = json.loads((ROOT / "release/v5/audits/manual-translations.json").read_bytes())
        self.assertEqual(set(audit["documents"]), set(LOCALES))
        for locale, report in audit["documents"].items():
            self.assertEqual(report["locale"], locale)
            self.assertEqual(report["accessibility_findings"], {"high": 0, "medium": 0, "low": 0})
            self.assertEqual(report["page_review"], "blocked_missing_bundled_soffice")
            self.assertEqual(report["artifact_sha256"], hashlib.sha256(self.outputs[locale].read_bytes()).hexdigest())


if __name__ == "__main__":
    unittest.main()
