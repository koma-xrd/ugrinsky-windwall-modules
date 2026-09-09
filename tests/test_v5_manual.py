"""Verify the released German manual at its DOCX and data boundaries.

Use the bundled document Python. The first authoring run requires the artifact
marker described in README; these tests subsequently rebuild into a temp folder.
"""

from copy import deepcopy
import hashlib
import importlib
import json
from pathlib import Path
import sys
import unittest
from zipfile import ZipFile

from tests.support import temporary_build_directory


ROOT = Path(__file__).resolve().parents[1]
DOCUMENT_RUNTIME = "codex-primary-runtime" in Path(sys.executable).parts
if DOCUMENT_RUNTIME:
    from docx import Document
    from docx.oxml.ns import qn
    from lxml import etree

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


@unittest.skipUnless(DOCUMENT_RUNTIME, "Run manual tests with the bundled document Python; see README")
class V5ManualTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.builder_path = ROOT / "scripts/manual/build_manual.py"
        if not cls.builder_path.is_file():
            return  # RED must never create a DOCX before the marker.
        cls.builder = importlib.import_module("scripts.manual.build_manual")
        cls.data_module = importlib.import_module("scripts.manual.manual_data")
        cls.temp = cls.enterClassContext(temporary_build_directory())
        cls.output = cls.temp / "manual.docx"
        cls.builder.build_v5_manual(ROOT, cls.output)
        cls.doc = Document(cls.output)
        with ZipFile(cls.output) as archive:
            cls.xml = etree.fromstring(archive.read("word/document.xml"))
        cls.text = "\n".join(cls.xml.xpath("//w:t/text()", namespaces=NS))
        cls.manifest = json.loads((ROOT / "release/v5/manifest.json").read_text(encoding="utf-8"))
        cls.figures = json.loads((ROOT / "release/v5/drawings/figures.json").read_text(encoding="utf-8"))

    def setUp(self):
        self.assertTrue(self.builder_path.is_file(), "V5 manual builder is absent")

    def test_a4_margins_language_and_title_metadata(self):
        for section in self.doc.sections:
            self.assertAlmostEqual(section.page_width.mm, 210, places=1)
            self.assertAlmostEqual(section.page_height.mm, 297, places=1)
            for side in ("top", "right", "bottom", "left"):
                self.assertAlmostEqual(getattr(section, side + "_margin").mm, 18, places=1)
        title = self.doc.paragraphs[0]
        self.assertEqual(title.style.name, "Title")
        self.assertEqual(title.text, "Ugrinsky Wind Wall V5 Bauanleitung")
        self.assertEqual(self.doc.core_properties.title, title.text)
        self.assertEqual(self.doc.core_properties.language, "de-DE")
        self.assertIn("V5", self.doc.core_properties.subject)
        self.assertEqual(self.doc.styles["Normal"].element.find(qn("w:rPr")).find(qn("w:lang")).get(qn("w:val")), "de-DE")
        self.assertFalse(self.xml.xpath("//w:pBdr | //w:pPr/w:shd | //w:txbxContent", namespaces=NS))

    def test_thirteen_logical_chapters_and_black_headings(self):
        headings = [p for p in self.doc.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual([int(p.text.split()[0]) for p in headings], list(range(1, 14)))
        for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
            style = self.doc.styles[name]
            self.assertEqual(str(style.font.color.rgb), "000000")
            self.assertFalse(style.element.xpath(".//w:pBdr"))
            color = style.element.xpath(".//w:color")[0]
            self.assertIsNone(color.get(qn("w:themeColor")))

    def test_every_current_figure_is_inline_captioned_and_has_source_alt_text(self):
        inline = self.xml.xpath("//wp:inline", namespaces=NS)
        self.assertEqual(len(inline), 15)
        self.assertFalse(self.xml.xpath("//wp:anchor", namespaces=NS))
        captions = [p.text for p in self.doc.paragraphs if p.style.name == "Caption"]
        for drawing in self.figures["figures"]:
            self.assertIn(f"{drawing['drawing_id']}  {drawing['caption']}", captions)
            matches = [item for item in inline if item.find(qn("wp:docPr")).get("descr") == drawing["alt_text"]]
            self.assertEqual(len(matches), 1)

    def test_tables_have_headers_borders_padding_and_no_fixed_heights(self):
        self.assertGreater(len(self.doc.tables), 5)
        for table in self.doc.tables:
            self.assertTrue(table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
            self.assertEqual(len(table._tbl.xpath("./w:tblPr/w:tblBorders/*[@w:color='D9D9D9']")), 6)
            self.assertTrue(table._tbl.xpath("./w:tblPr/w:tblCellMar"))
            self.assertFalse(table._tbl.xpath(".//w:trHeight[@w:hRule='exact']"))

    def test_bom_has_exact_inventory_quantities_and_release_filenames(self):
        bom = self.doc.tables[1]
        rows = {r.cells[2].text: r.cells[0].text for r in bom.rows[1:]}
        self.assertEqual(rows["base_rotor_module.stl"], "1")
        self.assertEqual(rows["standard_rotor_module.stl"], "5")
        self.assertEqual(len(rows), 8)
        self.assertEqual(sum(map(int, rows.values())), 12)
        for record in self.manifest["production_parts"] + self.manifest["coupons"]:
            self.assertIn(Path(record["stl_path"]).name, self.text)
            self.assertIn(Path(record["step_path"]).name, self.text)
        for assembly in self.manifest["assemblies"]:
            self.assertIn(assembly["step_path"], self.text)
        for label, count in (("M4 Deckelschrauben", "6"), ("Gefangene M4 Muttern", "6"),
                             ("Holzschrauben unten", "4"), ("Holzschrauben oben", "4"),
                             ("M8 Muttern", "3"), ("Magnete", "36")):
            records = [r for t in self.doc.tables for r in t.rows if label in [c.text for c in r.cells]]
            self.assertEqual(len(records), 1, label)
            self.assertEqual(records[0].cells[0].text, count)

    def test_current_bearings_ownership_order_retention_and_support(self):
        for required in ("25 × 42 × 11 mm", "8 × 22 × 7 mm", "Wellenscheibe dreht",
                         "Gehäusescheibe bleibt stationär", "Innenring", "Wälzkörper",
                         "unterhalb der stationären Spule", "Gehäuseschulter", "Verdrehsicherung",
                         "sechs M4", "vier Bodenlaschen", "von unten", "verlängerte M8",
                         "555,3 mm", "nicht wasserdicht"):
            self.assertIn(required, self.text)
        assembly = self.text.split("7 Generator montieren")[1].split("8 Rotorstapel montieren")[0]
        required_order = ("Unteren Magnetrotor einsetzen", "Spulenkassette einsetzen",
                          "Stationären Deckel montieren", "51105 einsetzen", "Base aufsetzen")
        positions = [assembly.index(value) for value in required_order]
        self.assertEqual(positions, sorted(positions))
        support = self.text.split("9 Am Holzrahmen montieren")[1].split("10 Spulen messen")[0]
        self.assertLess(support.index("M8-Klemmung"), support.index("Halter montieren"))

    def test_gap_polarity_print_and_electrical_experiment_limits(self):
        for value in ("1,5 mm", "bündig oder tiefer", "N S N S", "gegenüber", "PLA", "ASA",
                      "Bambu P2S", "0,4 mm", "Serpentinen", "0,18 mm", "20", "40", "80",
                      "gemessenen Drahtdurchmesser", "keine endgültige Windungszahl",
                      "nicht direkt an einen 48-V-Bleiakku", "Widerstand", "Drehzahl", "Temperatur",
                      "Zugentlastung", "Magnetrückhaltung", "Presspassung", "Dauerfestigkeit"):
            self.assertIn(value, self.text)
        matrix = next(t for t in self.doc.tables if t.cell(0, 0).text == "Draht")
        self.assertEqual([r.cells[1].text for r in matrix.rows[1:]], ["20", "40", "80"])
        for reference in ("Magnetfläche zur aktiven Wicklungsfläche", "Kassettenboden",
                          "Deckelmembran", "0,35 mm", "0,50 mm", "0,15 mm"):
            self.assertIn(reference, self.text)

    def test_no_obsolete_instructions_or_internal_placeholders(self):
        for forbidden in ("Top-Closure", "top_closure", "radiale Sicherungsschrauben",
                          "radialen Sicherungsschrauben", "provisional sleeve", "TODO", "TBD",
                          "PLACEHOLDER", "turn0search", "[[", "assets/manual/release", "V4.1"):
            self.assertNotIn(forbidden, self.text)
        self.assertIn("physisch nicht validiert", self.text)

    def test_build_is_byte_deterministic(self):
        other = self.temp / "second.docx"
        self.builder.build_v5_manual(ROOT, other)
        self.assertEqual(hashlib.sha256(self.output.read_bytes()).digest(), hashlib.sha256(other.read_bytes()).digest())
        released = ROOT / "release/v5/docs/Ugrinsky-Wind-Wall-V5-Bauanleitung.docx"
        self.assertEqual(self.output.read_bytes(), released.read_bytes(), "Rebuild the tracked V5 DOCX")

    def test_loader_rejects_stale_drawing_manifest_before_authoring(self):
        with temporary_build_directory() as folder:
            root = Path(folder)
            (root / "release/v5/drawings").mkdir(parents=True)
            (root / "release/v5/manifest.json").write_bytes((ROOT / "release/v5/manifest.json").read_bytes())
            figures = deepcopy(self.figures)
            figures["source_manifest_sha256"] = "0" * 64
            (root / "release/v5/drawings/figures.json").write_text(json.dumps(figures), encoding="utf-8")
            target = root / "unwritten.docx"
            with self.assertRaisesRegex(ValueError, "manifest hash"):
                self.builder.build_v5_manual(root, target)
            self.assertFalse(target.exists())

    def test_loader_rejects_missing_required_figure(self):
        with temporary_build_directory() as folder:
            root = Path(folder)
            (root / "release/v5/drawings").mkdir(parents=True)
            (root / "release/v5/manifest.json").write_bytes((ROOT / "release/v5/manifest.json").read_bytes())
            figures = deepcopy(self.figures)
            figures["figures"] = figures["figures"][:-1]
            (root / "release/v5/drawings/figures.json").write_text(json.dumps(figures), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "E01.*E15"):
                self.data_module.load_manual_data(root)


if __name__ == "__main__":
    unittest.main()
